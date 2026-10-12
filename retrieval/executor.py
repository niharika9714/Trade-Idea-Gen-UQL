"""Execute Step 5 plans against allowlisted SQLite and bounded original documents.

Metadata is the discovery surface; original sources are the evidence surface.
No LLM reasoning, generated claims, arbitrary SQL or outbound API calls.
"""
from __future__ import annotations

import calendar
import hashlib
import json
import re
import sqlite3
from datetime import date, datetime, timezone
from pathlib import Path
from typing import Any

from pydantic import BaseModel, Field
from app.config import DOCUMENTS_DIR, MARKET_DB_PATH, METADATA_DB_PATH
from core.enums import Authority, EvidenceType, SourceType
from core.models import Evidence, Provenance, RetrievalPlan, UQLTask


class AcquisitionResult(BaseModel):
    plan_id: str
    evidence: list[Evidence] = Field(default_factory=list)
    task_results: list[dict[str, Any]] = Field(default_factory=list)
    warnings: list[str] = Field(default_factory=list)
    deferred_agendas: list[str] = Field(default_factory=list)


def _ident(s: str) -> str:
    if not isinstance(s, str) or not re.fullmatch(r'[A-Za-z_][A-Za-z0-9_]*', s):
        raise ValueError('Unsafe SQL identifier')
    return '"' + s + '"'


def _catalog_file(metadata_db: Path, file_id: str, source_id: str) -> str:
    with sqlite3.connect(f'file:{metadata_db.resolve().as_posix()}?mode=ro', uri=True) as con:
        row = con.execute('SELECT relative_path FROM files WHERE file_id=? AND source_id=?', (file_id, source_id)).fetchone()
    if not row:
        raise ValueError('File locator not present in metadata catalog')
    return row[0]


def _catalog_columns(metadata_db: Path, source_id: str, table: str) -> set[str]:
    with sqlite3.connect(f'file:{metadata_db.resolve().as_posix()}?mode=ro', uri=True) as con:
        return {r[0] for r in con.execute('SELECT column_name FROM sql_columns WHERE source_id=? AND table_name=?', (source_id, table))}


def _date_range(scope: dict[str, Any]) -> tuple[str | None, str | None]:
    start, end = scope.get('start_date'), scope.get('end_date')
    if start and end:
        a, b = date.fromisoformat(str(start)), date.fromisoformat(str(end))
    elif scope.get('year') and scope.get('month'):
        months = {calendar.month_name[i].lower(): i for i in range(1, 13)}
        month = scope['month']
        m = months.get(str(month).lower(), int(month) if str(month).isdigit() else 0)
        y = int(scope['year'])
        if m not in range(1, 13):
            raise ValueError('Invalid month')
        a, b = date(y, m, 1), date(y, m, calendar.monthrange(y, m)[1])
    elif start or end:
        raise ValueError('Both start_date and end_date required')
    else:
        return None, None
    if a > b:
        raise ValueError('Start date after end date')
    return a.isoformat(), b.isoformat()


def _evidence(task: UQLTask, content: Any, locator: dict[str, Any], *, claim: str,
              authority: Authority, observation_time: datetime | None = None,
              publication_time: datetime | None = None) -> Evidence:
    seed = json.dumps([task.task_id, task.source_id, locator, content], sort_keys=True, default=str)
    eid = 'EV-' + hashlib.sha256(seed.encode('utf-8')).hexdigest()[:20]
    return Evidence(evidence_id=eid, agenda_id=task.agenda_id, evidence_type=EvidenceType.FACT,
                    claim=claim, content=content, authority=authority,
                    provenance=Provenance(source_id=task.source_id, source_type=task.source_type,
                                          locator=locator, observation_time=observation_time,
                                          publication_time=publication_time), confidence=1.0)


def _execute_sql(task: UQLTask, *, market_db: Path, metadata_db: Path,
                 max_rows: int, include_prior_close: bool) -> tuple[list[Evidence], list[str]]:
    locator = task.retrieval_query.get('locator', {})
    table = locator.get('table')
    if not table:
        raise ValueError('Missing SQL table locator')
    catalog_cols = _catalog_columns(metadata_db, task.source_id, table)
    if not catalog_cols:
        raise ValueError('SQL table not registered in metadata')
    requested = locator.get('columns', [])
    if not requested or not set(requested) <= catalog_cols:
        raise ValueError('Requested SQL columns not in catalog')
    filters = task.retrieval_query.get('filters', {})
    ticker = filters.get('ticker', {})
    values = ticker.get('values', []) if ticker else []
    if ticker and (ticker.get('operator') != 'IN' or ticker.get('column') != 'ticker' or 'ticker' not in catalog_cols or not values):
        raise ValueError('Unsupported ticker filter')
    if any(not isinstance(x, str) or len(x) > 30 for x in values) or len(values) > 25:
        raise ValueError('Invalid ticker values')
    scope = filters.get('temporal_scope', {})
    start, end = _date_range(scope)
    date_candidates = filters.get('date_column_candidates') or locator.get('date_columns', [])
    date_col = next((c for c in date_candidates if c in catalog_cols), None)
    if (start or end) and not date_col:
        raise ValueError('Temporal filter requested without a catalogued date column')
    clauses, params = [], []
    if values:
        clauses.append(f'"ticker" IN ({",".join("?" for _ in values)})')
        params.extend(values)
    if start:
        clauses.append(f'{_ident(date_col)} >= ? AND {_ident(date_col)} <= ?')
        params.extend([start, end])
    where = ' WHERE ' + ' AND '.join(clauses) if clauses else ''
    order = f' ORDER BY {_ident(date_col)} ASC' if date_col and start else (f' ORDER BY {_ident(date_col)} DESC' if date_col else '')
    query = f'SELECT {", ".join(_ident(c) for c in requested)} FROM {_ident(table)}{where}{order} LIMIT ?'
    if not market_db.is_file():
        raise FileNotFoundError(f'Market database not found: {market_db}')
    warnings = []
    with sqlite3.connect(f'file:{market_db.resolve().as_posix()}?mode=ro', uri=True, timeout=5) as con:
        con.row_factory = sqlite3.Row
        con.execute('PRAGMA query_only=ON')
        actual = {r[1] for r in con.execute(f'PRAGMA table_info({_ident(table)})')}
        if not set(requested) <= actual:
            raise ValueError('Market database schema differs from metadata catalog')
        rows = [dict(r) for r in con.execute(query, [*params, max_rows + 1]).fetchall()]
        if len(rows) > max_rows:
            warnings.append(f'{task.task_id}: SQL result exceeds max_rows={max_rows}; no incomplete series emitted')
            return [], warnings
        if not rows:
            return [], [f'{task.task_id}: no matching SQL rows']
        baseline = {}
        if include_prior_close and table == 'daily_prices' and start and values and 'close_price' in actual and date_col:
            for symbol in values:
                prior = con.execute(f'SELECT {_ident(date_col)}, "close_price" FROM {_ident(table)} '
                                    f'WHERE "ticker"=? AND {_ident(date_col)} < ? '
                                    f'ORDER BY {_ident(date_col)} DESC LIMIT 1', (symbol, start)).fetchone()
                if prior:
                    baseline[symbol] = {'date': prior[0], 'close_price': prior[1]}
    grouped: dict[str, list[dict[str, Any]]] = {}
    for row in rows:
        grouped.setdefault(str(row.get('ticker', 'ALL')), []).append(row)
    result = []
    for symbol, subset in grouped.items():
        evidence_locator = {'table': table, 'ticker': symbol if symbol != 'ALL' else None,
                            'date_column': date_col, 'start_date': start, 'end_date': end,
                            'row_count': len(subset), 'sql': query, 'parameters': [*params, max_rows + 1]}
        content = {'rows': subset, 'row_count': len(subset), 'complete': True,
                   'prior_period_close': baseline.get(symbol)}
        result.append(_evidence(task, content, evidence_locator,
                                claim=f'Observed {len(subset)} rows from {table} for {symbol}',
                                authority=Authority.INTERNAL_STRUCTURED))
    return result, warnings


def _pdf_pages(path: Path, max_units: int):
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    for i in range(min(len(reader.pages), max_units)):
        yield i + 1, reader.pages[i].extract_text() or ''
    return len(reader.pages) > max_units


def _docx_blocks(path: Path, max_units: int):
    from docx import Document
    doc = Document(str(path))
    # Preserve paragraph and table boundaries with stable document-order locators.
    count = 0
    for block in doc.element.body.iterchildren():
        if count >= max_units:
            return
        tag = block.tag.rsplit('}', 1)[-1]
        if tag == 'p':
            from docx.text.paragraph import Paragraph
            content = Paragraph(block, doc).text
        elif tag == 'tbl':
            from docx.table import Table
            tbl = Table(block, doc)
            content = '\n'.join(' | '.join(cell.text for cell in row.cells) for row in tbl.rows)
        else:
            continue
        count += 1
        yield count, content


def _pptx_slides(path: Path, max_units: int):
    from pptx import Presentation
    prs = Presentation(str(path))
    for i, slide in enumerate(prs.slides):
        if i >= max_units:
            return
        pieces = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                pieces.append(shape.text)
            if shape.has_table:
                pieces.extend(' | '.join(cell.text for cell in row.cells) for row in shape.table.rows)
            if shape.has_chart:
                pieces.append('[CHART: values require separate chart-data extraction]')
            if shape.shape_type == 13:
                pieces.append('[IMAGE: visual content not interpreted]')
        yield i + 1, '\n'.join(pieces)


def _msg_parts(path: Path, max_units: int):
    import extract_msg
    msg = extract_msg.Message(str(path))
    try:
        body = msg.body or ''
        yield 1, f'Subject: {msg.subject or ""}\nDate: {msg.date or ""}\n{body}'
        # Attachments intentionally not followed automatically: bounded, explicit retrieval only.
    finally:
        msg.close()


def _execute_file(task: UQLTask, *, documents_dir: Path, metadata_db: Path,
                  max_units: int, max_hits: int) -> tuple[list[Evidence], list[str]]:
    loc = task.retrieval_query.get('locator', {})
    fid = loc.get('file_id')
    if not fid:
        raise ValueError('Missing file_id')
    relative = _catalog_file(metadata_db, fid, task.source_id)
    root = documents_dir.resolve()
    path = (root / relative).resolve()
    if not path.is_relative_to(root) or not path.is_file():
        raise ValueError('Unsafe or missing source document path')
    suffix = path.suffix.lower()
    extractor = {'.pdf': _pdf_pages, '.docx': _docx_blocks, '.pptx': _pptx_slides, '.msg': _msg_parts}.get(suffix)
    if extractor is None:
        return [], [f'{task.task_id}: {suffix} requires a dedicated visual/format adapter; not extracted']
    terms = [str(t).strip().lower() for t in task.retrieval_query.get('query_terms', []) if len(str(t).strip()) >= 3]
    # Prioritize distinct longer terms; never match isolated generic one/two-character tokens.
    terms = sorted(set(terms), key=lambda x: (-len(x), x))
    ranked = []
    units_seen = 0
    for unit, content in extractor(path, max_units):
        units_seen += 1
        if not content.strip():
            continue
        lower = content.lower()
        hits = [t for t in terms if re.search(r'(?<!\w)' + re.escape(t) + r'(?!\w)', lower)]
        if hits:
            ranked.append((len(hits), unit, content, hits))
    ranked.sort(key=lambda x: (-x[0], x[1]))
    warnings = []
    if units_seen >= max_units:
        warnings.append(f'{task.task_id}: document inspection capped at {max_units} units; later content may be unsearched')
    if len(ranked) > max_hits:
        warnings.append(f'{task.task_id}: {len(ranked)} matching units; returning top {max_hits}')
    evidence = []
    for _, unit, content, hits in ranked[:max_hits]:
        # Return only bounded excerpts; don't imply images/charts have been understood.
        excerpt = content[:4000]
        locator = {'file_id': fid, 'relative_path': relative, 'unit': unit,
                   'unit_type': {'.pdf': 'page', '.pptx': 'slide', '.docx': 'block', '.msg': 'message'}[suffix],
                   'matched_terms': hits,
                   'citation': f'{relative}#page={unit}' if suffix == '.pdf' else f'{relative}#{"slide" if suffix == ".pptx" else "block"}={unit}'}
        evidence.append(_evidence(task, {'excerpt': excerpt, 'matched_terms': hits}, locator,
                                  claim=f'Source excerpt matched: {", ".join(hits[:5])}',
                                  authority=Authority.INTERNAL_COMMUNICATION if suffix == '.msg' else Authority.INTERNAL_RESEARCH))
    if not evidence:
        warnings.append(f'{task.task_id}: no matching text in inspected document units')
    return evidence, warnings


def acquire(plan: RetrievalPlan, *, metadata_db: str | Path = METADATA_DB_PATH,
            market_db: str | Path = MARKET_DB_PATH, documents_dir: str | Path = DOCUMENTS_DIR,
            max_sql_rows: int = 5000, max_document_units: int = 20,
            max_document_hits: int = 3, include_prior_close: bool = True) -> AcquisitionResult:
    if min(max_sql_rows, max_document_units, max_document_hits) < 1:
        raise ValueError('Execution limits must be positive')
    result = AcquisitionResult(plan_id=plan.plan_id)
    for task in plan.tasks:
        try:
            if task.source_type == SourceType.SQL:
                evidence, warnings = _execute_sql(task, market_db=Path(market_db),
                                                   metadata_db=Path(metadata_db), max_rows=max_sql_rows,
                                                   include_prior_close=include_prior_close)
            elif task.source_type in (SourceType.FILE, SourceType.EMAIL):
                evidence, warnings = _execute_file(task, documents_dir=Path(documents_dir),
                                                    metadata_db=Path(metadata_db), max_units=max_document_units,
                                                    max_hits=max_document_hits)
            else:
                evidence, warnings = [], [f'{task.task_id}: API execution not enabled; endpoint credentials and allowlist required']
            result.evidence.extend(evidence)
            result.warnings.extend(warnings)
            result.task_results.append({'task_id': task.task_id, 'status': 'COMPLETED' if evidence else 'NO_EVIDENCE',
                                        'evidence_count': len(evidence)})
        except Exception as exc:
            result.warnings.append(f'{task.task_id}: {type(exc).__name__}: {exc}')
            result.task_results.append({'task_id': task.task_id, 'status': 'FAILED', 'evidence_count': 0,
                                        'error': f'{type(exc).__name__}: {exc}'})
    return result
